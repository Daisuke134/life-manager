"""Render the sourced talk JSON to editable PPTX and a vector PDF.
Requires python-pptx/reportlab and the system Liberation Sans fonts.
This authors presentation artifacts; it does not implement Manager.
"""
from pathlib import Path
import argparse,json
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
FONT='/usr/share/fonts/truetype/liberation/'
pdfmetrics.registerFont(TTFont('TalkSans',FONT+'LiberationSans-Regular.ttf'))
pdfmetrics.registerFont(TTFont('TalkBold',FONT+'LiberationSans-Bold.ttf'))
W,H=960,540
NAVY='142C42';TEAL='137E7A';BG='F5F7F9';WHITE='FFFFFF';GRAY='506273';AMBER='946300'
def wrap(text,width,size,font='TalkSans'):
 lines=[]
 for raw in text.split('\n'):
  words=raw.split();line=''
  for word in words:
   trial=(line+' '+word).strip()
   if line and pdfmetrics.stringWidth(trial,font,size)>width:
    lines.append(line);line=word
   else:line=trial
  lines.append(line)
 return lines

def render(source,out):
 obj=json.loads(source.read_text());out.mkdir(parents=True,exist_ok=True)
 prs=Presentation();prs.slide_width=Pt(W);prs.slide_height=Pt(H)
 prs.core_properties.title=obj['title'];prs.core_properties.author=obj['author'];prs.core_properties.subject=obj['status']
 pdf=canvas.Canvas(str(out/'Human_Simulation_World_Models_Manager.pdf'),pagesize=(W,H));pdf.setTitle(obj['title']);pdf.setAuthor(obj['author'])
 def rect(s,x,y,w,h,fill,line=None):
  a=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h));a.fill.solid();a.fill.fore_color.rgb=RGBColor.from_string(fill);a.line.fill.background()
  pdf.setFillColor(HexColor('#'+fill));pdf.roundRect(x,H-y-h,w,h,8,fill=1,stroke=0)
 def text(s,lines,x,y,w,size=18,bold=False,color=NAVY,leading=None):
  leading=leading or size*1.22
  if isinstance(lines,str):lines=wrap(lines,w,size,'TalkBold' if bold else 'TalkSans')
  h=max(leading*len(lines)+6,leading+6)
  assert y+h<=H+1,(lines,y,h)
  a=s.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h));tf=a.text_frame;tf.word_wrap=False
  tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
  for i,t in enumerate(lines):
   p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=t;p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.line_spacing=Pt(leading);p.space_after=Pt(0)
   pdf.setFont('TalkBold' if bold else 'TalkSans',size);pdf.setFillColor(HexColor('#'+color));pdf.drawString(x,H-y-size-i*leading,t)
  return a
 for i,d in enumerate(obj['slides'],1):
  s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
  pdf.setFillColor(HexColor('#'+BG));pdf.rect(0,0,W,H,fill=1,stroke=0)
  text(s,d['kicker'],42,22,850,11,True,TEAL)
  titles=wrap(d['title'],870,30,'TalkBold');assert len(titles)<=2,d['title'];text(s,titles,42,50,876,30,True)
  text(s,d['subtitle'],42,135,876,17,False,GRAY)
  cards=d['cards'];n=len(cards)
  if n==4:geom=[(42+(j%2)*445,193+(j//2)*120,431,110) for j in range(4)];base=16
  else:geom=[(42+j*297,193,283,231) for j in range(n)];base=18
  for card,(x,y,w,h) in zip(cards,geom):
   rect(s,x,y,w,h,WHITE)
   heads=wrap(card['title'],w-28,17,'TalkBold');assert len(heads)<=2,card['title'];text(s,heads,x+14,y+12,w-28,17,True,TEAL)
   by=y+18+len(heads)*20.5;available=y+h-12-by
   size=base
   while True:
    body=wrap(card['body'],w-28,size)
    if len(body)*size*1.22+5<=available:break
    size-=.5
    assert size>=14,(i,card['title'],'text exceeds card')
   text(s,body,x+14,by,w-28,size)
  if d['callout']:
   text(s,d['callout'],42,449,876,14,True,AMBER)
  # Two lines of clickable source labels; full URLs are also in speaker notes.
  refs=d['sources']
  for row in range(2):
   group=refs[row*2:row*2+2]
   for j,r in enumerate(group):
    a=text(s,'Source: '+r['label'],42+j*450,487+row*14,440,9,False,GRAY)
    for p in a.text_frame.paragraphs:
     for run in p.runs:run.hyperlink.address=r['url']
    pdf.linkURL(r['url'],(42+j*450,H-(487+row*14)-12,482+j*450,H-(487+row*14)),relative=0)
  text(s,f'{i:02d} / {len(obj["slides"]):02d}',855,518,70,9,False,GRAY)
  notes=d['notes_ja']+'\n\nSources:\n'+'\n'.join(r['label']+'\n'+r['url'] for r in refs)
  s.notes_slide.notes_text_frame.text=notes
  pdf.showPage()
 prs.save(out/'Human_Simulation_World_Models_Manager.pptx');pdf.save()
 print(f'Rendered {len(obj["slides"])} slides to {out}; all card text fit checked')
if __name__=='__main__':
 a=argparse.ArgumentParser();a.add_argument('--source',type=Path,default=Path(__file__).with_name('2026-10-11-human-simulation-talk.json'));a.add_argument('--output',type=Path,required=True);args=a.parse_args();render(args.source,args.output)
