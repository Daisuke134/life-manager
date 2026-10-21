from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_THEME_COLOR
from pathlib import Path
import math, statistics

OUT=Path(__file__).resolve().parents[1] / 'output' / 'presentations' / 'Manager_Behavior_AGI_Research_Seminar_2026-10-21.pptx'\nOUT.parent.mkdir(parents=True, exist_ok=True)
HERO=Path(__file__).with_name('Manager_Research_Hero_2026-10-21.png')
if not HERO.is_file(): HERO=Path('/mnt/data/wide_cinematic_highly_detailed_illustrative_conc.png')
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
C={
'nv':(8,20,40),'dark':(16,37,58),'pane':(23,48,70),'pane2':(29,61,81),'border':(72,96,111),
'mint':(99,226,189),'teal':(26,166,144),'blue':(75,164,224),'ice':(233,250,252),'white':(252,253,254),
'muted':(171,193,206),'orange':(254,183,99),'red':(240,117,123),'violet':(179,153,251),'gray':(118,144,158),'yellow':(247,217,120)
}
def rgb(n): return RGBColor(*(C[n] if isinstance(n,str) else n))
def rect(sl,x,y,w,h,fill='pane',line=None,rounded=False):
    sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
    if line:sh.line.color.rgb=rgb(line);sh.line.width=Pt(1.1)
    else:sh.line.fill.background()
    if rounded: sh.adjustments[0]=.07
    return sh
def circle(sl,x,y,d,fill='teal',line=None):
    sh=sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d));sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
    if line: sh.line.color.rgb=rgb(line)
    else: sh.line.fill.background()
    return sh
def txt(sl,t,x,y,w,h,size=18,color='white',bold=False,align=None,margin=.04,font='Aptos',spacing=1.02):
    sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=Inches(margin);tf.margin_right=Inches(margin);tf.margin_top=Inches(.01);tf.margin_bottom=Inches(.01)
    for i,line in enumerate(t.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
        if align is not None:p.alignment=align
        p.line_spacing=spacing;p.space_before=Pt(0);p.space_after=Pt(1)
    return sh
def arr(sl,x1,y1,x2,y2,color='mint',width=2.5):
    sh=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=rgb(color);sh.line.width=Pt(width)
    return sh

def badge(sl,label,x,y,w,color='teal'):
    rect(sl,x,y,w,.36,color,rounded=True);txt(sl,label,x+.05,y+.052,w-.1,.23,10.5,'white',True,PP_ALIGN.CENTER)
def box(sl,title,body,x,y,w,h,accent='mint',bsize=16,tsize=19):
    rect(sl,x,y,w,h,'pane',rounded=True)
    rect(sl,x+.17,y+.24,.065,.37,accent)
    txt(sl,title,x+.36,y+.22,w-.5,.5,tsize,'white',True)
    txt(sl,body,x+.22,y+.86,w-.46,h-.96,bsize,'ice',False,spacing=1.12)
def base(title,sub='',tag='RESEARCH',source='',notes=''):
    s=prs.slides.add_slide(prs.slide_layouts[6]);rect(s,0,0,13.333,7.5,'nv')
    rect(s,0,0,13.333,.12,'teal');txt(s,'M A N A G E R   /   B E H A V I O R  ×  A G I',.68,.33,8.4,.24,11,'mint',True)
    badge(s,tag,11.3,.31,1.25,'pane2')
    txt(s,title,.71,.89,11.9,.63,23 if len(title)>65 else 25 if len(title)>53 else 29,'white',True)
    if sub:txt(s,sub,.73,1.70,11.85,.43,15,'muted')
    arr(s,.72,6.93,12.56,6.93,'border',1.1)
    txt(s,source[:175],.72,7.00,10.95,.25,8.6,'muted')
    txt(s,f'{len(prs.slides):02d}',12.14,6.98,.45,.25,11,'mint',True,PP_ALIGN.RIGHT)
    if notes:
        try:
            ns=s.notes_slide.notes_text_frame; ns.text='SPEAKER NOTES / SOURCES\n'+notes
        except Exception:
            pass
    return s

def section(title,sub,number,source=''):
    s=prs.slides.add_slide(prs.slide_layouts[6]);rect(s,0,0,13.333,7.5,'nv');rect(s,0,0,.15,7.5,'teal')
    txt(s,f'{number:02d}  /  MANAGER RESEARCH',.9,.8,9,.35,14,'mint',True)
    txt(s,title,.86,2.1,11.5,1.3,46,'white',True)
    txt(s,sub,.9,3.98,10.8,1.7,24,'ice')
    txt(s,source,.9,6.94,11.8,.26,10,'muted')
    return s

s=prs.slides.add_slide(prs.slide_layouts[6])
if HERO.is_file(): s.shapes.add_picture(str(HERO),0,0,width=prs.slide_width,height=prs.slide_height)
else: rect(s,0,0,13.333,7.5,'pane2')
rect(s,0,0,7.35,7.5,'nv');rect(s,.64,.85,.09,.59,'mint')
txt(s,'RESEARCH SEMINAR  ·  OCT 21 2026',.92,.92,5.6,.28,13,'mint',True)
txt(s,'From Human\nBehavior to\nBetter Worlds',.89,1.84,6.4,2.45,43,'white',True)
txt(s,'Behavior prediction · LLM agents · social simulation\nintervention science · transparent public benefit',.92,4.93,6.0,1.1,19,'ice')
txt(s,'MANAGER / ANICCA — END SUFFERING',.95,6.88,5.5,.28,12,'mint',True)
#1
s=base('Research question: can an AI improve the world—provably?',
'We are building a hypothesis, not claiming superintelligence.','OVERVIEW',notes='Distinguish ambition from achieved evidence. Manager remains research-stage. Net wellbeing changes not independently demonstrated.')
box(s,'SCIENTIFIC QUESTION','Can a public-information system select legitimate interventions that cause measurable net welfare improvements?',.78,2.45,7.4,3.4,'mint',22)
box(s,'EVIDENCE STANDARD','Prediction ≠ counterfactual ≠ execution ≠ impact.\n\nEvery step has its own benchmark.',8.43,2.45,4.05,3.4,'orange',18)
#2
s=base('The entire field in five layers','This is the simplest map from “synthetic people” to real-world benefit.','FOUNDATIONS')
for i,(a,b) in enumerate([('REPRESENT','Who is in the population?'),('PREDICT','Which actions or answers occur?'),('INTERACT','What emerges through social influence?'),('INTERVENE','What changes under a new policy?'),('VERIFY','Did the real world improve?')]):
    x=.72+i*2.55
    rect(s,x,2.38,2.36,2.45,'pane',rounded=True)
    circle(s,x+.91,2.62,.47,'teal' if i<3 else 'orange');txt(s,str(i+1),x+.93,2.68,.43,.23,13,'nv',True,PP_ALIGN.CENTER)
    txt(s,a,x+.18,3.24,2,.39,17,'white',True,PP_ALIGN.CENTER)
    txt(s,b,x+.18,3.85,2.0,.68,14,'muted',align=PP_ALIGN.CENTER)
    if i<4:arr(s,x+2.38,3.55,x+2.53,3.55,'mint',3)
txt(s,'A model can be convincing at layer 3 and still be wrong about layers 4–5.',.9,5.58,11.65,.7,21,'orange',True)
#3
s=base('Population synthesis is not behavior prediction','Synthetic population = a representation of composition, not a verified decision model.','FOUNDATIONS',notes='Synthetic population evaluation can compare margins and relationships. Matching age groups does not guarantee correct behavioral outcomes.')
box(s,'Population record','Age: 34\nHousehold: 3 people\nTransit access: low\nOccupation: service worker',.86,2.3,4.25,3.5,'blue',19)
box(s,'Missing behavioral mechanism','Will they travel?\nWill they switch jobs?\nWill a subsidy change their choice?\nWho influences them?',5.52,2.3,6.9,3.5,'orange',20)
#4
s=base('Behavior prediction: the conditional distribution','A statistical model maps a context to the probability of an outcome.','FOUNDATIONS',source='Behavioral forecasting principles; Centaur, Nature 2025 · nature.com/articles/s41586-025-09215-4')
txt(s,'P(Y = action | X = context)',.93,2.32,10.6,.75,34,'mint',True)
for i,(label,p) in enumerate([('WALK',.52),('BUS',.31),('CYCLE',.17)]):
    y=3.45+i*.74;txt(s,label,1.05,y,1.3,.36,16,'white',True)
    rect(s,2.53,y+.05,7.6,.29,'pane2',rounded=True);rect(s,2.53,y+.05,p*7.6,.29,'teal' if i==0 else 'blue' if i==1 else 'violet',rounded=True)
    txt(s,f'{int(p*100)}%',10.34,y-.005,1,.39,16,'ice',True)
txt(s,'This diagram uses fictional probabilities for teaching—not a measured dataset.',.99,6.24,11.3,.32,13,'orange')
#5
s=base('Three distinct accuracy questions','A model must be judged on the claim being made.','EVALUATION',source='SocSci210 · arxiv.org/abs/2509.05830 · Artificial Societies Benchmark · arxiv.org/abs/2609.30030')
box(s,'Individual accuracy','Did it predict a particular response?',.75,2.39,3.78,2.6,'blue',17)
box(s,'Population TVD','Did the distribution of responses match observed humans?',4.79,2.39,3.78,2.6,'mint',17)
box(s,'Causal fidelity','Did it predict the change induced by treatment?',8.84,2.39,3.78,2.6,'orange',17)
txt(s,'These three questions can produce three different winners.',1.03,5.75,11.25,.54,21,'white',True)
#6
s=base('TVD, Brier score, and calibration','A transparent forecast must communicate uncertainty.','MATH',source='TVD = 1/2 Σ |P_i − Q_i|; Brier = mean((p_i − y_i)^2)')
box(s,'Total variation distance','TVD(P,Q) = ½ Σ |Pᵢ − Qᵢ|\n\nHow far apart are predicted and observed distributions?',.78,2.31,5.68,3.2,'mint',20)
box(s,'Brier score','Brier = average (p − y)²\n\nPenalizes incorrect confidence, not just wrong labels.',6.77,2.31,5.68,3.2,'blue',20)
txt(s,'Calibration: among predictions made at 70% confidence, roughly 70% should occur.',.95,6.05,11.75,.51,16,'orange')
#7
s=base('Prediction versus counterfactual simulation','A what-if statement is not an observed causal effect.','CAUSAL INFERENCE')
rect(s,.81,2.37,5.65,2.88,'pane',rounded=True)
txt(s,'OBSERVED',1.1,2.68,4.8,.42,22,'mint',True)
txt(s,'P(Y | X)\nWhat do people tend to do in context X?',1.1,3.35,4.85,1.19,23,'white')
rect(s,6.77,2.37,5.65,2.88,'pane',rounded=True)
txt(s,'INTERVENTIONAL',7.05,2.68,4.8,.42,22,'orange',True)
txt(s,'P(Y | do(T = 1))\nWhat changes if we intervene?',7.05,3.35,4.9,1.19,23,'white')
txt(s,'Valid effect estimates need assignment logic, a comparator, bias controls and uncertainty.',.98,5.85,11.7,.61,17,'ice')
#8
s=base('Causal graphs expose mistaken explanations','Illustrative DAG: context can cause both treatment selection and outcome.','CAUSAL INFERENCE',source='DoWhy · github.com/py-why/dowhy · HeartSteps original trial PMC6401341')
for cx,cy,lab,color in [(1.18,3.19,'CONTEXT','violet'),(5.16,3.19,'TREATMENT','blue'),(9.36,3.19,'OUTCOME','mint')]:
    rect(s,cx,cy,2.55,.90,'pane',rounded=True);txt(s,lab,cx+.1,cy+.31,2.35,.4,18,color,True,PP_ALIGN.CENTER)
arr(s,3.75,3.65,5.15,3.65,'mint',3);arr(s,7.72,3.65,9.36,3.65,'mint',3)
arr(s,2.6,3.15,10.0,2.48,'orange',1.9)
txt(s,'Confounding',4.9,2.25,2.3,.33,15,'orange',True)
txt(s,'Randomized assignment weakens the context → treatment selection problem.',1.14,5.02,11.3,.65,20,'white')
#9
s=base('Experimental design: never skip the comparator','Micro-randomization allows many intervention opportunities per person.','EXPERIMENTS',source='HeartSteps · Klasnja et al. 2019 · PMC6401341')
for i,(title,body) in enumerate([('ELIGIBLE','Could a decision be randomized?'),('RANDOMIZE','Send suggestion vs do not send.'),('FOLLOW UP','Measure behavior in the next 30 minutes.'),('ESTIMATE','Cluster adjustment, missingness, harms.')]):
    x=.84+i*3.06
    box(s,title,body,x,2.4,2.82,2.75,'mint' if i<3 else 'orange',15,17)
    if i<3:arr(s,x+2.84,3.55,x+3.04,3.55,'mint',3)
txt(s,'The same participant contributes repeated observations: rows are not IID.',.95,5.85,11.4,.62,18,'orange')
#10
s=section('How did the field get here?','Cognitive modeling → generative agents → networked societies → foundation models → causal evaluation.',1,'Selected primary research 1996–2026; see full references at end.')
#11
s=base('Scientific lineage: six milestones','A research map, not a ranking of the “best” model.','HISTORY')
events=[('1996','Sugarscape','bottom-up artificial society'),('2006–10','Social influence','network experiments'),('2023','Generative Agents','memory / reflection / plan'),('2023–26','Concordia','entity + Game Master'),('2025','Centaur / Socrates','behavior fine-tuning'),('2026','Validity debates','causal and subgroup checks')]
for i,(year,name,desc) in enumerate(events):
    col=i%3;row=i//3;x=.82+col*4.10;y=2.14+row*1.89
    rect(s,x,y,3.84,1.66,'pane',rounded=True)
    txt(s,year,x+.22,y+.20,3.25,.32,15,'orange',True)
    txt(s,name,x+.22,y+.63,3.31,.37,18,'white',True)
    txt(s,desc,x+.22,y+1.17,3.40,.3,12,'muted')
#12
s=base('Generative Agents: believable daily lives','The 2023 architecture connected memory, reflection and planning.','RESEARCH',source='Park et al., Generative Agents (2023) · arxiv.org/abs/2304.03442')
for i,(a,b) in enumerate([('MEMORY','Store observations'),('REFLECTION','Summarize recurring themes'),('PLANNING','Select future actions'),('WORLD','Events change context')]):
    x=.95+i*3.05;box(s,a,b,x,2.48,2.80,2.60,'mint' if i<3 else 'violet',16,18)
    if i<3:arr(s,x+2.83,3.65,x+3.02,3.65,'mint',2.5)
txt(s,'Believable behavior in a simulated town is not a measured real-world welfare effect.',.98,5.73,11.6,.72,18,'orange',True)
#13
s=base('Concordia: the actor–world separation','An agent proposes; the Game Master adjudicates simulated outcomes.','FRAMEWORK',source='Google DeepMind Concordia official · github.com/google-deepmind/concordia')
box(s,'AGENT A','Goal + memory + action policy',.94,2.41,3.42,2.12,'blue',16)
box(s,'GAME MASTER','World rules, context and transition',4.99,2.41,3.43,2.12,'mint',16)
box(s,'AGENT B','Another agent reacts and learns',9.03,2.41,3.42,2.12,'blue',16)
arr(s,4.39,3.40,4.96,3.40,'mint',3);arr(s,8.45,3.40,9.0,3.40,'mint',3)
txt(s,'Game Master ≠ permission to change real facilities, people or finances.',1.07,5.47,11.2,.59,20,'orange',True)
#14
s=base('Centaur: language-formatted cognition tasks','A foundation model of human cognition trained on Psych-101.','RESEARCH',source='Binz et al., Nature (2025) · nature.com/articles/s41586-025-09215-4')
for x,v,l in [(.95,'160','experimental paradigms'),(4.95,'60,092','participants'),(8.95,'10.68M','observed choices')]:
    rect(s,x,2.30,3.45,1.83,'pane',rounded=True);txt(s,v,x+.12,2.72,3.1,.56,30,'mint',True,PP_ALIGN.CENTER);txt(s,l,x+.10,3.47,3.12,.42,16,'muted',align=PP_ALIGN.CENTER)
txt(s,'Generalization evaluated for held-out participants, cover stories, task structures and domains.',1.02,4.95,11.2,1.11,22,'ice')
#15
s=base('Socrates / SocSci210: social-science responses','Separate unseen studies from unseen people within known studies.','RESEARCH',source='Kolluri et al. (2025) · arxiv.org/abs/2509.05830 · github.com/akaashkolluri/socrates')
for i,(big,sub) in enumerate([('210','social-science studies'),('170 / 40','seen / unseen study IDs'),('~2.9M','reported response records')]):
    x=.87+i*4.11;rect(s,x,2.47,3.86,1.82,'pane',rounded=True);txt(s,big,x+.17,2.80,3.48,.62,30,'mint',True,PP_ALIGN.CENTER);txt(s,sub,x+.14,3.57,3.54,.41,14,'muted',align=PP_ALIGN.CENTER)
txt(s,'Project status: study-ID split verified; real 40-study human benchmark still not independently run by Manager.',.96,5.13,11.4,1.08,20,'orange')
#16
s=base('Research instruments: Expected Parrot / EDSL','Think survey × scenario × model × respondent × result.','FRAMEWORK',source='Expected Parrot · docs.expectedparrot.com · github.com/expectedparrot/edsl')
for i,(a,b) in enumerate([('QUESTION','What must be answered?'),('SCENARIO','What stimulus is shown?'),('AGENT','Who is modeled?'),('MODEL','Which predictor replies?'),('RESULT','How do we score vs humans?')]):
    x=.72+i*2.54;rect(s,x,2.46,2.34,2.38,'pane',rounded=True);txt(s,a,x+.08,3.00,2.17,.41,16,'mint',True,PP_ALIGN.CENTER);txt(s,b,x+.2,3.73,1.96,.8,14,'ice',align=PP_ALIGN.CENTER)
txt(s,'Reproducible tools enable experiments; human validity must still be demonstrated.',1.00,5.67,11.4,.6,19,'orange')
#17
s=base('Population-scale systems: AgentTorch, OASIS, AgentSociety','Different ways to model thousands of agents and their interactions.','FRAMEWORK')
box(s,'AgentTorch','Differentiable agent-based frameworks; inspect calibration and numerical assumptions.',.82,2.30,3.80,3.35,'blue',17)
box(s,'OASIS','LLM agents interact on modeled social platforms; examine collective dynamics.',4.81,2.30,3.8,3.35,'mint',17)
box(s,'AgentSociety','Connect social environments with economic and behavioral state.',8.81,2.30,3.8,3.35,'violet',17)
txt(s,'Sources: github.com/AgentTorch/AgentTorch · github.com/camel-ai/oasis · github.com/tsinghua-fib-lab/AgentSociety',.92,6.60,11.50,.33,10,'muted')
#18
s=section('Who is building what?','Commercial research platforms and academic simulation libraries solve different parts of the problem.',2)
#19
s=base('Simile: scenario testing for behavioral decisions','Company-described prediction and human-observation validation.','COMPANY',source='simile.com/blog/confidence · simile.com (review original technical description before quoting performance)')
box(s,'INPUT','Audience / transactions / interview or survey evidence',.79,2.37,3.66,2.64,'blue',17)
box(s,'SIMULATION','Change price, product, message or conditions',4.84,2.37,3.66,2.64,'mint',17)
box(s,'OUTPUT + CONFIDENCE','Expected response plus uncertainty / evaluation',8.88,2.37,3.66,2.64,'orange',17)
txt(s,'Counterfactual questions inside a model are useful, but different from executed real-world interventions.',.95,5.64,11.7,.73,17,'ice')
#20
s=base('Aaru: published survey + purchase validation','Aaru reports measured error on human survey answers and an observed card-purchase panel.','COMPANY',source='Aaru Research, Sep 24 2026 · aaru.com/publications/population-simulation-at-the-replication-floor-research')
for i,(x,tag,big,small) in enumerate([(0,'Survey questions','2,993','Company evaluation'),(1,'Unadjusted mean TVD','7.62 pts','Different scale and benchmark'),(2,'Observed purchase data','CARD PANEL','Behavioral ranking comparison')]):
    a=.81+i*4.08;rect(s,a,2.37,3.8,2.6,'pane',rounded=True);txt(s,tag,a+.19,2.70,3.47,.36,16,'mint',True);txt(s,big,a+.20,3.36,3.40,.60,29,'white',True);txt(s,small,a+.19,4.18,3.38,.51,13,'muted')
txt(s,'Vendor’s own reported evaluation; not a universal accuracy guarantee and not Manager-run.',1.03,5.64,11.4,.63,17,'orange')
#21
s=base('Artificial Societies / Radiant: the product','Synthetic audience research with an explicit social network.','COMPANY',source='Official product · societies.ai/radiant · accessed 21 Oct 2026',notes='Explicit distinction: 38% per-ticket displayed on product is illustrative demo synthetic survey, not direct human survey. Not all reported benchmark claims are independently reproduced.')
for i,(a,b) in enumerate([('SOCIETIES','Connected persona cohorts'),('SURVEYS','Quantitative distributions'),('INTERVIEWS','Persona-level follow-up'),('FOCUS GROUPS','Interaction and opinion shifts'),('ASSISTANT','Cross-study synthesis')]):
    x=.72+i*2.55;rect(s,x,2.41,2.37,2.74,'pane',rounded=True);txt(s,a,x+.13,2.88,2.08,.45,16,'mint',True,PP_ALIGN.CENTER);txt(s,b,x+.16,3.75,2.06,.88,15,'ice',align=PP_ALIGN.CENTER)
txt(s,'Radiant product demo numbers are synthetic output—not observed survey measurements.',.95,5.91,11.55,.67,17,'orange')
#22
s=base('Radiant: from observed data to networked audiences','A proposed architecture reconstructed from the company’s public descriptions.','COMPANY',source='societies.ai/knowledge/the-science-behind-artificial-societies · /radiant')
for i,(a,b) in enumerate([('GROUNDING','Public / licensed observations'),('PERSONA GENESIS','Inferred beliefs & attitudes'),('SOCIETY ASSEMBLY','Relationships and influence'),('SIMULATION','Scenario, responses, discussion'),('VALIDATION','Humans, subgroups, effects')]):
    x=.71+i*2.55;rect(s,x,2.43,2.39,2.62,'pane',rounded=True);txt(s,a,x+.12,2.89,2.11,.5,15,'mint',True,PP_ALIGN.CENTER);txt(s,b,x+.13,3.66,2.13,.87,15,'ice',align=PP_ALIGN.CENTER)
    if i<4:arr(s,x+2.39,3.73,x+2.54,3.73,'mint',3)
txt(s,'Exact production code, learned belief representation and network update laws are proprietary / not verified.',.92,5.82,11.65,.80,17,'orange')
#23
s=base('Radiant: important validation work—and its limits','Internal, construct, and external validity are not interchangeable.','BENCHMARK',source='Chidichimo et al. 2026, arxiv.org/abs/2609.30030 · societies.ai/as-benchmark')
box(s,'INTERNAL VALIDITY','Test-retest, coherent beliefs, paraphrases, survey design effects.',.81,2.38,3.80,2.82,'blue',17)
box(s,'CONSTRUCT VALIDITY','Trait covariance, convergent and discriminant relationships.',4.77,2.38,3.80,2.82,'mint',17)
box(s,'EXTERNAL VALIDITY','Human distributions, subgroups, treatment effects and structure.',8.73,2.38,3.80,2.82,'orange',17)
txt(s,'Research preprint proposes 11 tests. Company web pages also reference an eight-test vendor guide; these scopes differ.',1.01,5.60,11.25,.86,16,'ice')
#24
s=base('Radiant headline scores are not general-purpose causal proof','Always preserve what was scored, on which dataset, and by whom.','BENCHMARK',source='societies.ai/as-benchmark · Artificial Societies internal benchmark, 1 Oct 2026')
for i,(tag,val,desc) in enumerate([('Reported mean TVD','0.08','Artificial Societies model'),('Best generic LLM','0.31','Vendor’s comparison on that benchmark'),('Responses evaluated','68,280','Across seven human survey sources')]):
    x=.85+i*4.1;rect(s,x,2.32,3.80,2.17,'pane',rounded=True);txt(s,tag,x+.2,2.63,3.46,.36,15,'mint',True);txt(s,val,x+.2,3.11,3.43,.72,33,'white',True);txt(s,desc,x+.2,4.08,3.42,.31,12,'muted')
txt(s,'These are vendor-reported survey evaluation results, not independently demonstrated social welfare gains.',1.01,5.33,11.2,1.03,20,'orange')
#25
s=base('The ecosystem matrix: avoid a false winner','Different research objectives require different metrics and experimental designs.','SYNTHESIS',source='Sources: Simile, Aaru, Radiant, EDSL, Centaur, SocSci210, Concordia official technical reports')
headers=['APPROACH','EXAMPLE','PRIMARY QUESTION','CRITICAL TEST']
rows=[
 ['Cognitive model','Centaur','Choices across tasks','Held-out tasks'],
 ['Survey model','Socrates','Experimental responses','Unseen-study errors'],
 ['Audience simulation','Simile / Aaru / Radiant','Scenario responses','Human panel vs output'],
 ['Agent environments','Concordia / OASIS','Social interactions','Mechanism & empirical fit'],
 ['Research instruments','Expected Parrot','Comparable study execution','Reproducibility'],
 ['Impact operations','Manager — proposed','Independently positive net welfare','Authorized, causal readback']
]
widths=[2.45,2.85,3.35,3.7]; xs=[.70,3.18,6.04,9.45];
for j,h in enumerate(headers):
    rect(s,xs[j],2.22,widths[j],.52,'pane2');txt(s,h,xs[j]+.1,2.37,widths[j]-.15,.22,12,'mint',True)
for i,r in enumerate(rows):
    y=2.76+i*.53;bg='pane' if i%2==0 else 'dark'
    for j,v in enumerate(r):
        rect(s,xs[j],y,widths[j],.51,bg);txt(s,v,xs[j]+.09,y+.14,widths[j]-.16,.33,11.5,'white' if i<5 else 'orange',i==5)
#26
s=section('What has Manager actually built?','From honest replication failures to a verified prediction pipeline—not yet field impact.',3)
#27
s=base('Phase 03: HeartSteps and reproducibility','Exact replication failure is a scientific result, not a success badge.','MANAGER-RUN',source='Manager pinned CI · PMC6401341 original trial · project reproducibility report')
for x,v,t in [(.88,'8,274','Released decision rows'),(4.88,'6,061','Original available points'),(8.88,'6,330','Different exploratory subset')]:
    rect(s,x,2.44,3.48,2.15,'pane',rounded=True);txt(s,v,x+.15,2.95,3.12,.65,35,'mint',True,PP_ALIGN.CENTER);txt(s,t,x+.1,3.93,3.27,.42,14,'muted',align=PP_ALIGN.CENTER)
txt(s,'The original 7,540/6,061 analytic selection could not be reconstructed from the public export.',.94,5.35,11.65,.64,18,'orange')
#28
s=base('Phase 04: genuine observed aggregate forecasting','2011 train → 2012 holdout, on UCI daily bike-rental counts.','MANAGER-RUN',source='Original UCI dataset · DOI 10.24432/C5W894 · Manager CI run 39545601251')
for i,(a,b) in enumerate([('TRAIN','365 observed days (2011)'),('HOLD OUT','366 days (2012)'),('RESULT','MAE 2489 global vs 2304 season'),('UNCERTAINTY','95% gain interval crosses zero')]):
    x=.79+i*3.06;box(s,a,b,x,2.40,2.82,3.01,'mint' if i<3 else 'orange',16,18)
txt(s,'Measured aggregate mobility demand—not SocSci210 response prediction and not human causal intervention proof.',.99,5.79,11.4,.77,17,'orange')
#29
s=base('Why predictive fit does not establish a good intervention','A causal counterexample built into Manager tests.','MANAGER-RUN',source='Manager synthetic methods tests; conceptual example only—not human experimental evidence')
rect(s,.82,2.41,5.78,2.6,'pane',rounded=True);txt(s,'MODEL MATCHES MARGINAL',1.06,2.79,5.22,.38,19,'mint',True);txt(s,'P(Y = 1) = 0.5\nTVD = 0.0',1.07,3.54,4.8,1.04,26,'white',True)
rect(s,6.77,2.41,5.78,2.6,'pane',rounded=True);txt(s,'BUT MISSES TREATMENT',7.04,2.79,5.16,.38,19,'orange',True);txt(s,'True constructed ATE = 0.4\nBlind model ATE = 0.0',7.05,3.54,4.95,1.04,24,'white',True)
txt(s,'This is an artificial mathematical example—not an estimated effect on a real population.',.98,5.75,11.6,.81,18,'orange')
#30
s=base('Manager: an evidence-constrained public-benefit kernel','Reusable architecture across finance, access, health and environments.','PROPOSAL',source='Manager project architecture; proposed or prototype components, not verified deployed effects')
for i,(a,b) in enumerate([('OBSERVE','Public sources & provenance'),('HYPOTHESIZE','Causal candidates'),('COMPARE','Predict + NO_ACTION'),('AUTHORIZE','Standing grants & guardrails'),('EXECUTE','Idempotent narrow adapters'),('VERIFY','Independent external readback'),('LEARN','Holdout, rollback, audit')]):
    row=i//4;col=i%4;x=.80+col*3.09;y=2.24+row*1.9
    box(s,a,b,x,y,2.86,1.64,'mint' if i<3 else 'orange' if i in (3,4) else 'blue',12,16)
txt(s,'No action is a valid choice. An agent may not create its own legal permissions.',.95,6.43,11.8,.4,16,'orange')
#31
s=base('Transparent Blessings: open an evidence receipt','A public change requires more than a self-reported “success.”','PROPOSAL')
checks=[('ISSUE','Source, baseline, affected system'),('OPTIONS','Comparator, NO_ACTION, forecasts'),('AUTHORITY','Grant, cap, scope, revocation'),('EXECUTION','External receipt / uncertainty'),('MEASUREMENT','Independent dated change'),('IMPACT','Causal net effect and harms'),('CHALLENGE','Appeal, correction and rollback')]
for i,(a,b) in enumerate(checks):
    y=2.16+i*.61;rect(s,.9,y,11.63,.54,'pane',rounded=True);txt(s,f'{i+1:02d}   {a}',1.1,y+.16,2.9,.27,13,'mint',True);txt(s,b,4.15,y+.16,8.05,.27,14,'white')
txt(s,'Current verified Manager beneficiary improvements: NOT MEASURED.',1.00,6.60,11.45,.26,12,'orange',True)
#32
s=base('Testable thesis: what makes Manager new?','Do not market an unproved difference as a demonstrated advantage.','HYPOTHESES')
box(s,'H1 · BETTER CHOICE','Evidence-linked causal tests improve intervention selection versus simple baselines.',.81,2.33,3.79,3.31,'mint',17)
box(s,'H2 · SAFER AUTONOMY','Standing rights/budget gates reduce unauthorized or unverified effects.',4.78,2.33,3.79,3.31,'blue',17)
box(s,'H3 · VERIFIED BENEFIT','Independent measurement shows net impact across a second real domain.',8.75,2.33,3.79,3.31,'orange',17)
txt(s,'Each hypothesis must have a preregistered metric, threshold and null/failure report.',1.00,5.97,11.37,.6,17,'ice')
#33
s=base('Research proposal: three experiments worth doing','Small, falsifiable studies before large promises.','NEXT WORK')
box(s,'EXP 1 · BEHAVIOR','Licensed held-out studies: frequency baseline vs open fine-tuned model; subgroup and treatment-error metrics.',.78,2.28,3.84,3.45,'mint',16)
box(s,'EXP 2 · SOCIAL','Networked vs isolated agent responses on held-out human group-influence experiments.',4.76,2.28,3.83,3.45,'blue',16)
box(s,'EXP 3 · FIELD','An authorized public-information intervention with independent comparator, net welfare and harm accounting.',8.74,2.28,3.83,3.45,'orange',16)
txt(s,'Success criteria, permission scope and evaluation datasets must be fixed before looking at the outcomes.',.94,6.16,11.57,.59,15,'ice')
#34
s=base('What remains scientifically unproven','Clear limits are a research strength.','LIMITATIONS')
for i,(a,b) in enumerate([('REAL HUMAN BENCHMARK','SocSci210 study-outcome processing and scope of permission not independently read back.'),('CAUSAL FIDELITY','Simulation intervention effects have not been shown accurate across external RCTs.'),('REAL ACTION','No audited world-changing Manager executor has verified positive field outcomes.'),('GENERALITY','No replicated benefit across finance, health, environment, employment.'),('GOVERNANCE','Legitimate, contestable authority and transparent harm accounting remain design obligations.')]):
    y=2.25+i*.85;rect(s,.8,y,11.73,.72,'pane',rounded=True);txt(s,a,1.02,y+.20,3.15,.33,14,'orange',True);txt(s,b,4.14,y+.19,8.0,.42,14,'ice')
#35
s=section('Mission: End suffering.','The ambition is to make the world better—and to publish evidence that anyone can inspect.',4,'Research proposal, not a statement of achieved superintelligence or field impact.')
#36 references
s=base('Primary sources I — cognition and simulation','Open the source before accepting the summary.','SOURCES')
refs=[
('Generative Agents · Park et al. (2023)','https://arxiv.org/abs/2304.03442'),
('Centaur / Psych-101 · Binz et al. (2025)','https://www.nature.com/articles/s41586-025-09215-4'),
('Socrates / SocSci210 · Kolluri et al. (2025)','https://arxiv.org/abs/2509.05830'),
('Concordia · Google DeepMind','https://github.com/google-deepmind/concordia'),
('HeartSteps RCT (2019)','https://pmc.ncbi.nlm.nih.gov/articles/PMC6401341/'),
('Expected Parrot EDSL','https://github.com/expectedparrot/edsl'),
('AgentTorch / OASIS / AgentSociety','https://github.com/AgentTorch/AgentTorch | github.com/camel-ai/oasis')]
for i,(a,b) in enumerate(refs):
    y=2.20+i*.58;txt(s,a,.9,y,5.5,.26,13,'mint',True);txt(s,b,6.3,y,6.1,.27,12,'ice')
#37 references
s=base('Primary sources II — companies and validity','Vendor performance reports ≠ independent external replication.','SOURCES')
refs=[
('Radiant official product','https://societies.ai/radiant'),
('Artificial Societies method','https://societies.ai/knowledge/the-science-behind-artificial-societies'),
('Artificial Societies benchmark 2026','https://arxiv.org/abs/2609.30030'),
('AI chatbot homophily · He et al.','https://doi.org/10.1111/bjop.12764'),
('Aaru evaluation (company-reported)','https://aaru.com/publications/population-simulation-at-the-replication-floor-research'),
('Simile method and uncertainty','https://www.simile.com/blog/confidence'),
('UCI Bike Sharing dataset','https://doi.org/10.24432/C5W894')]
for i,(a,b) in enumerate(refs):
    y=2.20+i*.58;txt(s,a,.9,y,5.55,.26,13,'mint',True);txt(s,b,6.25,y,6.15,.27,12,'ice')
#38 closing
s=prs.slides.add_slide(prs.slide_layouts[6]);rect(s,0,0,13.333,7.5,'nv');rect(s,.85,1.37,.10,2.05,'mint')
txt(s,'THE RESEARCH COMMITMENT',1.21,1.47,10.6,.36,15,'mint',True)
txt(s,'Predict carefully.\nIntervene legitimately.\nVerify independently.',1.16,2.18,11.3,2.42,37,'white',True)
txt(s,'Transparent Blessings are a hypothesis to prove—not a number to invent.',1.21,5.72,11.2,.94,20,'ice')

prs.save(OUT)
print('SAVED',OUT,'slides',len(prs.slides),'bytes',OUT.stat().st_size)