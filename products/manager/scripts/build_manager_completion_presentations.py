"""Build the three Y007 decks from their reviewed source JSON.

The requested repo-native python-pptx workflow reuses the 16:9 size and navy,
mint and orange palette of build_behavior_agi_lab_seminar_v3.py. Content and
flat evidence layouts are new. Existing decks and the original builder are
never overwritten. All visible text, tables and speaker notes stay editable.
"""

from __future__ import annotations

import json
import os
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/presentations/MANAGER_COMPLETION_DECK_SOURCES_20261022.json"
OUT = ROOT / "docs/presentations/artifacts"
WIDTH, HEIGHT = 13.333, 7.5
FONT = "DejaVu Sans"
COLOR = {
    "background": "081428",
    "header": "173046",
    "alternate": "10253A",
    "white": "FCFDFE",
    "text": "E9FAFC",
    "muted": "ABC1CE",
    "mint": "63E2BD",
    "orange": "FEB763",
}
SHORT_SOURCE = {
    "M-SPEC": "Manager spec",
    "M-RELEASE": "Release record",
    "M-CI": "Readback CI",
    "M-PAGES": "Pages run",
    "M-PUBLIC": "Public prototype",
    "M-DEMO": "Captured decision JSON",
    "M-VIDEO": "Demo video",
    "M-WEDGE": "Provisional wedge",
    "M-MAP": "Measurement map",
    "M-PEOPLE": "Researcher map",
    "M-PITCH": "Pitch evidence",
    "HEARTSTEPS-PAPER": "HeartSteps",
    "JITAI-PAPER": "JITAI",
    "MEGASTUDY-PAPER": "Milkman et al. 2021",
    "CASH-TRIAL": "Haushofer & Shapiro 2016",
    "GENERATIVE-PAPER": "Park et al. 2023",
    "ONS-LONELINESS": "ONS loneliness guidance",
    "WHO5-INSTRUMENT": "WHO-5 2024",
    "WHODAS": "WHO WHODAS 2.0",
    "CFPB-SCORECARD": "CFPB scorecard",
    "CLS-2024-25": "Community Life Survey 2024/25",
    "MHLW-SUPPORT": "MHLW support overview",
    "GOVUK-COMPLETION": "GOV.UK completion guidance",
    "SIMILE-OFFICIAL": "Simile official site",
}


def rgb(name: str) -> RGBColor:
    return RGBColor.from_string(COLOR[name])


def text(slide, value, x, y, w, h, size=22, color="text", bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0)
    tf.margin_top = tf.margin_bottom = Inches(0.025)
    for i, line in enumerate(value.split("\n")):
        paragraph = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        paragraph.text = line
        paragraph.font.name = FONT
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = rgb(color)
        paragraph.line_spacing = 1.08
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(4)
    return shape


def notes(slide, spec, sources):
    lines = ["日本語補助ノート", "", spec["ja"], "", "CLAIM TYPE", spec["claim_type"], "", "PRIMARY SOURCES"]
    for source_id in spec["sources"]:
        source = sources[source_id]
        lines.extend([source["label"], source.get("url", source.get("path", "")), ""])
    slide.notes_slide.notes_text_frame.text = "\n".join(lines)


def footer(slide, number, spec, sources):
    shape = text(slide, "", .77, 7.00, 11.25, .31, 9.5, "muted")
    paragraph = shape.text_frame.paragraphs[0]
    prefix = paragraph.add_run()
    prefix.text = "Sources: "
    for i, source_id in enumerate(spec["sources"]):
        if i:
            paragraph.add_run().text = " / "
        run = paragraph.add_run()
        run.text = SHORT_SOURCE.get(source_id, source_id)
        source = sources[source_id]
        if source.get("url"):
            run.hyperlink.address = source["url"]
        run.font.color.rgb = rgb("muted")
    text(slide, f"{number:02d}", 12.1, 6.98, .45, .32, 11, "mint", True)


def draw_table(slide, spec):
    count = len(spec["rows"])
    height = 3.65 if count < 6 else 3.84
    y = 2.28
    shape = slide.shapes.add_table(count + 1, len(spec["headers"]), Inches(.77), Inches(y), Inches(11.85), Inches(height))
    table = shape.table
    for column, width in zip(table.columns, spec["widths"]):
        column.width = Inches(width)
    table.rows[0].height = Inches(.55)
    for row in list(table.rows)[1:]:
        row.height = Inches((height - .55) / count)
    font_size = 17.5 if count >= 6 else 18
    for r, values in enumerate([spec["headers"], *spec["rows"]]):
        for c, value in enumerate(values):
            cell = table.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(.13)
            cell.margin_top = cell.margin_bottom = Inches(.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb("header" if r == 0 else "alternate" if r % 2 else "background")
            cell.text = value
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(17 if r == 0 else font_size)
                p.font.bold = r == 0
                p.font.color.rgb = rgb("mint" if r == 0 else "text")
                p.line_spacing = 1.08
                p.space_after = Pt(0)


def build_slide(prs, spec, number, sources):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("background")
    kind = spec["type"]
    if kind == "cover":
        text(slide, spec["subtitle"], .79, .83, 11.75, .65, 20, "mint")
        text(slide, spec["title"], .75, 2.07, 11.8, 2.10, 46, "white", True)
        text(slide, spec["statement"], .79, 5.1, 11.5, 1.0, 24, "text")
    else:
        text(slide, spec["title"], .75, .63, 11.9, 1.03, 31, "white", True)
        if spec.get("subtitle"):
            text(slide, spec["subtitle"], .79, 1.65, 11.7, .52, 16, "muted")
        if kind == "table":
            draw_table(slide, spec)
        elif kind == "metrics":
            for metric, x in zip(spec["metrics"], [.80, 7.00]):
                text(slide, metric["value"], x, 2.45, 5.2, 1.25, 74, "mint", True)
                text(slide, metric["label"], x, 4.03, 5.25, .83, 23, "white", True)
                text(slide, metric["detail"], x, 5.13, 5.25, .65, 16, "muted")
        elif kind == "split":
            for side, x in [("left", .80), ("right", 7.00)]:
                text(slide, spec[f"{side}_title"], x, 2.42, 5.2, .65, 23, "mint", True)
                text(slide, spec[f"{side}_body"], x, 3.24, 5.2, 2.72, 21, "text")
        else:
            raise ValueError(f"Unknown slide type: {kind}")
        if spec.get("statement"):
            text(slide, spec["statement"], .80, 6.28, 11.8, .55, 17.5, "orange")
    notes(slide, spec, sources)
    footer(slide, number, spec, sources)


def source_markdown(source, parent):
    if source.get("url"):
        target = source["url"]
    else:
        path = source["path"]
        if path.startswith("products/manager/"):
            target = os.path.relpath(ROOT / path[len("products/manager/"):], parent)
        else:
            target = path
    return f'[{source["label"]}]({target})'


def markdown(deck, data, target):
    out = [f'# Manager completion presentation: {deck["audience"]}', "",
           "English slide source and Japanese speaker notes. Prepared 2026-10-22 JST / 2026-10-21 UTC. No scheduled presentation date is asserted.", "",
           f'Editable deck: [PowerPoint](artifacts/{deck["name"]}.pptx). Visible slides: English. Full Japanese support and exact sources follow each slide.', "",
           "The P2 claim concerns an existing public digital artifact. Human outcomes remain NOT_MEASURED. Forecasts, proposed studies and observed results are explicitly distinguished.", ""]
    for i, spec in enumerate(deck["slides"], 1):
        out.extend([f'## Slide {i}: {spec["title"].replace(chr(10), " ")}', "", f'**Claim type:** `{spec["claim_type"]}`', ""])
        if spec.get("subtitle"):
            out.extend([spec["subtitle"], ""])
        if spec["type"] == "table":
            out.append("| " + " | ".join(spec["headers"]) + " |")
            out.append("| " + " | ".join("---" for _ in spec["headers"]) + " |")
            out.extend("| " + " | ".join(row) + " |" for row in spec["rows"])
            out.append("")
        elif spec["type"] == "metrics":
            out.extend(["| Value | Meaning | Scope |", "|---|---|---|"])
            out.extend(f'| {m["value"]} | {m["label"]} | {m["detail"]} |' for m in spec["metrics"])
            out.append("")
        elif spec["type"] == "split":
            for side in ["left", "right"]:
                out.extend([f'### {spec[f"{side}_title"]}', "", spec[f"{side}_body"], ""])
        if spec.get("statement"):
            out.extend([spec["statement"], ""])
        out.extend(["### 日本語補助ノート", "", spec["ja"], "", "### Sources", ""])
        out.extend(f'- {source_markdown(data["sources"][s], target.parent)}' for s in spec["sources"])
        out.append("")
    return "\n".join(out)


def main():
    data = json.loads(SOURCE.read_text())
    OUT.mkdir(parents=True, exist_ok=True)
    report = []
    for deck in data["decks"]:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(WIDTH), Inches(HEIGHT)
        # Office renders hyperlink colors from the theme, even when a run has
        # an explicit color. Set accessible link colors for the dark canvas.
        namespace = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        for part in prs.part.package.iter_parts():
            if str(part.partname).startswith("/ppt/theme/"):
                theme = etree.fromstring(part.blob)
                for name, color in [("hlink", "muted"), ("folHlink", "mint")]:
                    for node in theme.xpath(f"//a:clrScheme/a:{name}", namespaces=namespace):
                        for child in list(node):
                            node.remove(child)
                        value = etree.SubElement(node, f'{{{namespace["a"]}}}srgbClr')
                        value.set("val", COLOR[color])
                part._blob = etree.tostring(theme, xml_declaration=True, encoding="UTF-8", standalone=True)
        prs.core_properties.title = f'Manager completion: {deck["audience"]}'
        prs.core_properties.subject = "English slides and Japanese notes with primary sources"
        prs.core_properties.author = "Manager project"
        for number, spec in enumerate(deck["slides"], 1):
            for source_id in spec["sources"]:
                if source_id not in data["sources"]:
                    raise ValueError(f"Unknown source: {source_id}")
            build_slide(prs, spec, number, data["sources"])
        path = OUT / f'{deck["name"]}.pptx'
        prs.save(path)
        md_path = ROOT / f'docs/presentations/MANAGER_COMPLETION_{deck["key"]}_20261022.md'
        md_path.write_text(markdown(deck, data, md_path))
        report.append({"deck": str(path.relative_to(ROOT)), "slides": len(prs.slides), "bytes": path.stat().st_size, "markdown": str(md_path.relative_to(ROOT))})
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
